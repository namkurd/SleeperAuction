#!/usr/bin/env python3
"""ONE-TIME script: import the 2013-2020 standings from the `Lifetime` tab of DTF Auction.xlsx.

    pip install openpyxl
    python scripts/build_standings.py "DTF Auction.xlsx"

Writes data/history_standings.csv (season, manager, rank, rumbles, h2h_w, h2h_l, pf, teams).
`rank` is the Rumble R column: the finish by Rumbles (official standings since 2022; before that the rank
each team would have had). 2021 on is worked out from Sleeper by update_data.py with the same rule, and
matches the sheet's Rumble R for every team in 2021-2025, so only 2013-2020 are imported here.
The Lifetime tab has no 2012 block, so 2012 has no standings.
"""
import csv
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
LAST_SHEET_SEASON = 2020


def main(path):
    aliases = json.loads((ROOT / "config" / "aliases.json").read_text(encoding="utf-8"))["owner_aliases"]
    known = {}
    with open(ROOT / "data" / "history_2012_2020.csv", newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            known.setdefault(int(r["season"]), set()).add(r["manager"])
    ws = openpyxl.load_workbook(path, data_only=True)["Lifetime"]
    head = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, str):
            head.setdefault(v, c)
    need = ["Rumbles", "Rumble R", "H2H W", "H2H L", "PF"]
    col = {k: head[k] for k in need}
    out, r = [], 1
    while r <= ws.max_row:
        season = ws.cell(r, 1).value
        if not isinstance(season, int) or season > LAST_SHEET_SEASON:
            r += 1
            continue
        block, rr = [], r + 1
        while rr <= ws.max_row and ws.cell(rr, 1).value:
            name = ws.cell(rr, 1).value
            m = aliases.get(name, name)
            if m not in known.get(season, set()):
                sys.exit(f"{season}: '{name}' in the Lifetime tab is not a manager in the {season} draft. "
                         f"Add it to owner_aliases in config/aliases.json.")
            block.append({"season": season, "manager": m, "rank": ws.cell(rr, col["Rumble R"]).value,
                          "rumbles": ws.cell(rr, col["Rumbles"]).value, "h2h_w": ws.cell(rr, col["H2H W"]).value,
                          "h2h_l": ws.cell(rr, col["H2H L"]).value, "pf": round(ws.cell(rr, col["PF"]).value)})
            rr += 1
        for b in block:      # ties share a rank (competition ranking: 1, 2, 2, 4, ...)
            if b["rank"] != 1 + sum(1 for o in block if o["rumbles"] > b["rumbles"]):
                sys.exit(f"{season}: {b['manager']}'s Rumble R ({b['rank']}) does not follow from the Rumbles column")
        for b in block:
            b["teams"] = len(block)
        out += block
        r = rr
    with open(ROOT / "data" / "history_standings.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["season", "manager", "rank", "rumbles", "h2h_w", "h2h_l", "pf", "teams"],
                           lineterminator="\n")
        w.writeheader()
        w.writerows(out)
    print(f"wrote {len(out)} rows, seasons {out[0]['season']}-{out[-1]['season']}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "DTF Auction.xlsx")
