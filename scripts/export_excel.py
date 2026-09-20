#!/usr/bin/env python3
"""Write the whole archive back out as an Excel workbook laid out like the original DTF Auction file.

    pip install openpyxl
    python scripts/export_excel.py                 # writes DTF Auction - Updated.xlsx in the repo root
    python scripts/export_excel.py some/path.xlsx

Sheets
  Read me          what is in the workbook and how to paste it into the original
  Drafting         same seven columns as the original `Drafting` tab (Owner, Year, Roster, Depth_Chart,
                   Position, Player, Price): corrected names, prices and positions, lineup slots by price
  Auctions         the boards, styled after the original `Auctions` tab (short names like "C. McCaffrey")
  Player origins   NFL team and college of every pick
  Name fixes       every spreadsheet name that changed (and the few left as typed), and how
  Differences      where the old spreadsheet and Sleeper disagree (2021 on)

It reads data/picks.csv, so run scripts/update_data.py first if you want the newest season in it.
Not part of the automatic update.
"""
import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
FONT = "Arial"
SUFFIXES = {"jr", "sr", "ii", "iii", "iv", "v"}


def read_csv(name):
    with open(ROOT / "data" / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def norm(name):
    s = "".join(c for c in unicodedata.normalize("NFKD", str(name)) if not unicodedata.combining(c)).lower()
    s = re.sub(r"[.'’`\-]", "", s)
    return " ".join(t for t in re.sub(r"[^a-z0-9 ]", " ", s).split() if t not in SUFFIXES)


def short_name(name, pos):
    if pos == "DEF":
        return name
    t = name.strip().split()
    if len(t) < 2 or re.fullmatch(r"[A-Za-z]\.", t[0]):
        return name
    return t[0][0].upper() + ". " + " ".join(t[1:])


def mix(a, b, t):
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x * (1 - t) + y * t):02X}" for x, y in zip(ca, cb))


def luminance(hexcolor):
    def f(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hexcolor[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def main(out_path):
    picks = read_csv("picks.csv")
    for r in picks:
        r["season"], r["price"], r["row"], r["depth"] = int(r["season"]), int(r["price"]), int(r["row"]), int(r["depth"])
    site = json.loads((ROOT / "site" / "data.json").read_text(encoding="utf-8"))
    boards = site["boards"]
    seasons = sorted({r["season"] for r in picks}, reverse=True)

    by_team = defaultdict(list)
    for r in picks:
        by_team[(r["season"], r["manager"])].append(r)

    def roster_label(slot):
        return "D/ST" if slot in ("D/ST", "DEF") else slot

    # ordered like the boards: newest season first, managers left to right, lines top to bottom
    ordered = []
    for s in seasons:
        for m in boards[str(s)]["order"]:
            ordered += sorted(by_team[(s, m)], key=lambda r: r["row"])

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(name=FONT, bold=True, color="FFFFFF")
    base = Font(name=FONT, size=10)

    # ------------------------------------------------------------------ Read me
    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("DTF Club auction archive - corrected data", "title"),
        ("", None),
        ("What is here", "h"),
        ("Drafting: the same seven columns as your original Drafting tab (A:G), rebuilt from the cleaned archive. Paste it over A1:G in the original (2,492 rows, same as before) and refresh the pivot table.", None),
        ("Auctions: every draft board, styled after your Auctions tab. Delete the old Auctions tab and move this one in, or paste over it.", None),
        ("Player origins, Name fixes, Differences: extra reference tabs. They are not needed by the original workbook. Name fixes lists every name that changed, plus 22 names that could not be matched to a Sleeper player and were left as typed (check those by hand).", None),
        ("", None),
        ("What changed compared with the original", "h"),
        ("Names: nicknames, misspellings and initials are replaced by the real player (see Name fixes). Older years now use full names, e.g. cmc -> Christian McCaffrey.", None),
        ("Prices and positions: 2021 on come from Sleeper, the record of the draft itself. They match your sheet except where listed under Differences (two $1 prices in 2021 and a few positions mislabelled in the sheet).", None),
        ("Depth_Chart: strictly by price. A team's priciest player at a position is 1 (QB1, RB1, ...), the next is 2, and so on; ties go to the earlier pick.", None),
        ("Roster (lineup slot), the same rule for every team and year: priciest QB in QB; 2nd QB in SFLEX (2023 on); two priciest RBs and two priciest WRs in the RB and WR slots; next-best RB or WR in FLEX (2020-21 had two FLEX, 2022 a WR/TE flex); priciest TE; K; D/ST; everyone else BE, priciest first.", None),
        ("Owner: co-owners are folded into one name (Fernando/Ben = Ben, Jake/David = Jake, Oleg/Eric = Oleg) and Haan appears as Rohaan, the same names the website uses.", None),
        ("Roster shows D/ST for defenses in every year (the original mixed DEF, DST and D/ST). Position still says DEF.", None),
        ("Player origins gives each pick's NFL team (that season's roster; a mid-season trade can show either team) and college.", None),
        ("", None),
        ("Why a separate file", "h"),
        ("Your original has a pivot table, slicers and charts. Re-saving it with a script would strip the slicers, so this workbook holds only the data for you to paste.", None),
    ]
    for i, (text, kind) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, size=14 if kind == "title" else 11, bold=kind in ("title", "h"))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    # ------------------------------------------------------------------ Drafting
    ws = wb.create_sheet("Drafting")
    cols = ["Owner", "Year", "Roster", "Depth_Chart", "Position", "Player", "Price"]
    ws.append(cols)
    for r in ordered:
        ws.append([r["manager"], r["season"], roster_label(r["slot"]), f"{r['position']}{r['depth']}",
                   r["position"], r["player"], r["price"]])
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = base
    for col, w in zip("ABCDEFG", (14, 8, 9, 13, 10, 28, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    # ------------------------------------------------------------------ Auctions (boards)
    ws = wb.create_sheet("Auctions")
    peach = PatternFill("solid", fgColor="F9CB9C")
    thin = Side(style="thin", color="7F8FA6")
    row = 1
    max_cols = 1
    for s in seasons:
        order, slots = boards[str(s)]["order"], boards[str(s)]["slots"]
        n = len(order)
        max_cols = max(max_cols, 1 + 2 * n)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1 + 2 * n)
        c = ws.cell(row=row, column=1, value=f"{s} Auction Draft")
        c.font, c.fill, c.alignment = Font(name=FONT, size=14, bold=True), peach, Alignment(horizontal="center")
        for k in range(2, 2 + 2 * n):
            ws.cell(row=row, column=k).fill = peach
        # shade: darkest at both ends, lightest at the 5th column (as in the original)
        for i, m in enumerate(order):
            p = min(4, n - 1)
            d = (p - i) / p if i <= p and p else ((i - p) / (n - 1 - p) if n - 1 - p else 0)
            fill_hex = mix("DBE8F8", "0F3D75", min(1, d * 0.96))
            fg = "0B1C33" if luminance(fill_hex) > 0.32 else "FFFFFF"
            fill = PatternFill("solid", fgColor=fill_hex)
            c0 = 2 + 2 * i
            ws.merge_cells(start_row=row + 1, start_column=c0, end_row=row + 1, end_column=c0 + 1)
            h = ws.cell(row=row + 1, column=c0, value=m)
            h.font, h.alignment = Font(name=FONT, size=10, bold=True, color=fg), Alignment(horizontal="center")
            ws.cell(row=row + 1, column=c0 + 1).fill = fill
            h.fill = fill
            for j, label in enumerate(("Player", "Price")):
                x = ws.cell(row=row + 2, column=c0 + j, value=label)
                x.font, x.fill = Font(name=FONT, size=9, bold=True, color=fg), fill
                x.alignment = Alignment(horizontal="center" if j == 0 else "right")
            team = {r["row"]: r for r in by_team[(s, m)]}
            for k, slot in enumerate(slots):
                pr = team.get(k)
                pc = ws.cell(row=row + 3 + k, column=c0, value=short_name(pr["player"], pr["position"]) if pr else None)
                pp = ws.cell(row=row + 3 + k, column=c0 + 1, value=pr["price"] if pr else None)
                pc.font = pp.font = Font(name=FONT, size=10, bold=True, color=fg)
                pc.fill = pp.fill = fill
                pp.number_format = '"$"#,##0'
                pp.alignment = Alignment(horizontal="right")
                pc.border = Border(bottom=thin)
                pp.border = Border(bottom=thin)
        navy = PatternFill("solid", fgColor="0F3563")
        for k in range(3):
            ws.cell(row=row + 1 + k, column=1).fill = navy
        for k, slot in enumerate(slots):
            a = ws.cell(row=row + 3 + k, column=1, value=slot)
            a.font, a.fill, a.alignment = Font(name=FONT, size=10, bold=True, color="FFFFFF"), navy, Alignment(horizontal="center")
        row += 3 + len(slots) + 1                               # one blank row between years
    ws.column_dimensions["A"].width = 8
    for k in range(2, max_cols + 1):
        ws.column_dimensions[get_column_letter(k)].width = 17 if k % 2 == 0 else 6.5
    ws.sheet_view.showGridLines = False

    # ------------------------------------------------------------------ Player origins
    ws = wb.create_sheet("Player origins")
    ws.append(["Year", "Owner", "Player", "Position", "NFL team", "College"])
    for r in ordered:
        ws.append([r["season"], r["manager"], r["player"], r["position"], r["team"] or "Unknown",
                   "" if r["position"] == "DEF" else (r["college"] or "Unknown")])
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
    for row_cells in ws.iter_rows(min_row=2):
        for c in row_cells:
            c.font = base
    for col, w in zip("ABCDEF", (8, 14, 28, 10, 26, 24)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # ------------------------------------------------------------------ Name fixes
    ws = wb.create_sheet("Name fixes")
    ws.append(["Year", "As typed in the old sheet", "Now shown as", "Position", "Price", "How it was matched"])
    how = {"override": "Nickname or confirmation from the alias list",
           "override+posmismatch": "Alias list (the sheet had the wrong position)",
           "override-unmatched": "Alias list (player not in Sleeper's database, so no player id)",
           "fuzzy": "Name similarity, checked against the player's age in that year",
           "fuzzy+posmismatch": "Name similarity (the sheet had the wrong position)",
           "fuzzy+age": "Name similarity; two players share the name, chose the one whose age fits",
           "fuzzy+rank": "Name similarity; several candidates, chose the fantasy-relevant one",
           "unresolved": "Could not be matched; kept as typed"}
    changed = []
    for r in read_csv("history_name_map.csv"):
        if r["method"] == "team":
            continue
        if norm(r["name_raw"]) != norm(r["resolved"]) or r["method"] == "unresolved":
            changed.append((int(r["season"]), r["name_raw"], r["resolved"], r["position"], int(float(r["price"])),
                            how.get(r["method"], r["method"])))
    for row_vals in sorted(changed, key=lambda t: (t[0], t[1].lower())):
        ws.append(list(row_vals))
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
    for row_cells in ws.iter_rows(min_row=2):
        for c in row_cells:
            c.font = base
    for col, w in zip("ABCDEF", (8, 26, 28, 10, 8, 62)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"

    # ------------------------------------------------------------------ Differences
    ws = wb.create_sheet("Differences")
    ws.append(["Season", "Manager", "What differs (old spreadsheet vs Sleeper; Sleeper is used)"])
    rep = (ROOT / "data" / "reconcile_report.md").read_text(encoding="utf-8").splitlines()
    for line in rep:
        m = re.match(r"- (\d{4}) ([^:]+): (.*)", line)
        if m:
            ws.append([int(m.group(1)), m.group(2), m.group(3)])
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
    for row_cells in ws.iter_rows(min_row=2):
        for c in row_cells:
            c.font = base
    for col, w in zip("ABC", (8, 14, 110)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    wb.save(out_path)
    print(f"wrote {out_path}: {len(ordered)} picks, {len(seasons)} seasons, {len(changed)} name fixes")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "DTF Auction - Updated.xlsx"))
